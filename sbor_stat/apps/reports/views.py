from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.decorators import api_view
from apps.users.models import User
from django.utils.dateparse import parse_date
from django.db.models import Prefetch
from datetime import datetime
from .models import Report
from .serializers import ReportSerializer
import calendar
from rest_framework import status
import io
import openpyxl
from django.http import HttpResponse
from datetime import date, timedelta
from dateutil.relativedelta import relativedelta


@api_view(['POST'])
def reports_by_pck(request):
    pck = request.data.get('pck')
    if not pck:
        return Response({'error': 'ПЦК не указано'}, status=400)

    users_in_pck = User.objects.filter(pck=pck, role='teacher')
    reports = Report.objects.filter(user__in=users_in_pck)
    serializer = ReportSerializer(reports, many=True)
    return Response(serializer.data)
    

@api_view(['POST'])
def check_report(request, report_id):
    try:
        report = Report.objects.get(id=report_id)
    except Report.DoesNotExist:
        return Response({'error': 'Отчет не найден'}, status=404)

    report.checked_by_head = True
    report.save()

    return Response({'success': 'Отчет отмечен как проверенный'})


@api_view(['POST'])
def filtered_reports_by_pck(request):
    pck = request.data.get('pck')
    month = int(request.data.get('month')) 
    year = int(request.data.get('year'))

    if not pck or not month or not year:
        return Response({'error': 'Недостаточно данных'}, status=400)

    start = (date(year, month, 20) - relativedelta(months=1)).replace(day=20)
    end = date(year, month, 20)

    users_in_pck = User.objects.filter(pck=pck, role='teacher')
    reports = Report.objects.filter(user__in=users_in_pck, date__gte=start, date__lt=end)
    serializer = ReportSerializer(reports, many=True)
    return Response(serializer.data)


@api_view(['POST'])
def export_reports_excel(request):
    pck = request.data.get('pck')
    month = int(request.data.get('month'))
    year = int(request.data.get('year'))
    teacher_id = request.data.get('teacher_id')
    checked_only = request.data.get('checked_only', False)

    if not pck or not month or not year:
        return Response({'error': 'Недостаточно данных'}, status=400)

    start = (date(year, month, 20) - relativedelta(months=1)).replace(day=20)
    end = date(year, month, 20)

    users_in_pck = User.objects.filter(pck=pck, role='teacher')
    if teacher_id:
        users_in_pck = users_in_pck.filter(id=int(teacher_id))

    reports = Report.objects.filter(
        user__in=users_in_pck,
        date__gte=start,
        date__lt=end
    ).prefetch_related('criteria', 'user')

    if checked_only in [True, 'true', 'True']:
        reports = reports.filter(checked_by_head=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Отчеты"

    ws.append([
        "Преподаватель",
        "Мероприятие",
        "Дата",
        "Результат",
        "Проверено",
        "Критерии",
        "Баллы (суммарно)"
    ])

    total_points = 0

    for report in reports:
        criteria_list = report.criteria.all()
        criteria_names = ', '.join([c.name for c in criteria_list])
        criteria_points = sum(c.point for c in criteria_list)
        total_points += criteria_points

        ws.append([
            report.user.username,
            report.event_name,
            report.date.strftime('%Y-%m-%d'),
            report.result,
            "Да" if report.checked_by_head else "Нет",
            criteria_names,
            criteria_points
        ])

    ws.append([""] * 6 + [total_points])
    ws.cell(row=ws.max_row, column=6).value = "ИТОГО"

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f"Отчеты_{month}_{year}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


@api_view(['POST'])
def create_report(request):
    try:
        user = User.objects.get(username=request.data.get('username'))
    except User.DoesNotExist:
        return Response({'error': 'Пользователь не найден'}, status=404)

    serializer = ReportSerializer(data=request.data)
    if serializer.is_valid():
        serializer.save(user=user)
        return Response(serializer.data, status=201)
    return Response(serializer.errors, status=400)


@api_view(['POST'])
def user_reports(request):
    username = request.data.get('username')
    if not username:
        return Response({'error': 'Имя пользователя не указано'}, status=400)

    try:
        user = User.objects.get(username=username)
    except User.DoesNotExist:
        return Response({'error': 'Пользователь не найден'}, status=404)

    today = date.today()
    start_range = (today - relativedelta(months=1)).replace(day=20)
    end_range = today.replace(day=20)

    reports = Report.objects.filter(user=user)
    serialized = ReportSerializer(reports, many=True)

    # помечаем, какие можно редактировать
    for r in serialized.data:
        r_date = next((rep.date for rep in reports if rep.id == r['id']), None)
        r['editable'] = start_range <= r_date <= end_range

    return Response(serialized.data)


@api_view(['GET'])
def report_detail(request, pk):
    try:
        report = Report.objects.get(pk=pk)
        serializer = ReportSerializer(report)
        return Response(serializer.data)
    except Report.DoesNotExist:
        return Response({'error': 'Отчет не найден'}, status=status.HTTP_404_NOT_FOUND)
    

@api_view(['PUT'])
def update_report(request, report_id):
    try:
        report = Report.objects.get(id=report_id)
    except Report.DoesNotExist:
        return Response({'error': 'Отчет не найден'}, status=404)

    user_id = request.data.get('user_id') 
    role = request.data.get('role')

    if not user_id or not role:
        return Response({'error': 'Нет данных пользователя'}, status=400)

    if role == 'teacher':
        if report.user.id != int(user_id):
            return Response({'error': 'Нельзя редактировать чужой отчет'}, status=403)
        if report.checked_by_head:
            return Response({'error': 'Отчет уже проверен и не может быть изменен'}, status=403)

    serializer = ReportSerializer(report, data=request.data, partial=True)
    if serializer.is_valid():
        serializer.save()
        return Response(serializer.data)
    return Response(serializer.errors, status=400)


@api_view(['POST'])
def get_admin_reports(request):
    pck = request.data.get('pck')

    if not pck:
        return Response({'error': 'ПЦК не указано'}, status=400)

    users = User.objects.filter(pck=pck, role='teacher')
    reports = Report.objects.filter(user__in=users, checked_by_head=True).prefetch_related('criteria')

    data = []
    for r in reports:
        data.append({
            'id': r.id,
            'user_id': r.user.id,
            'user': r.user.username,
            'event_name': r.event_name,
            'about_event': r.about_event,
            'members': r.members,
            'result': r.result,
            'proofs': r.proofs,
            'checked_by_head': r.checked_by_head,
            'date': r.date.strftime('%Y-%m-%d'),
            'criteria': [{'id': c.id, 'name': c.name, 'point': c.point} for c in r.criteria.all()],
        })

    return Response(data)


@api_view(['POST'])
def export_admin_reports_excel(request):
    pck = request.data.get('pck')
    if not pck:
        return Response({'error': 'ПЦК не указано'}, status=400)

    users = User.objects.filter(pck=pck, role='teacher')
    reports = Report.objects.filter(user__in=users, checked_by_head=True).prefetch_related('criteria')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Отчеты"

    ws.append([
        "Преподаватель",
        "Мероприятие",
        "Дата",
        "Результат",
        "Критерии",
        "Баллы"
    ])

    total_points = 0

    for r in reports:
        criteria_list = r.criteria.all()
        crit_names = ', '.join([c.name for c in criteria_list])
        crit_points = sum(c.point for c in criteria_list)
        total_points += crit_points

        ws.append([
            r.user.username,
            r.event_name,
            r.date.strftime('%Y-%m-%d'),
            r.result,
            crit_names,
            crit_points
        ])

    ws.append([""] * 5 + [total_points])
    ws.cell(row=ws.max_row, column=5).value = "ИТОГО"

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="admin_reports_{pck}.xlsx"'
    wb.save(response)
    return response

